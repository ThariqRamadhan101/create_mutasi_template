from openpyxl import load_workbook
from collections import OrderedDict
import re
from generate_from_data import (
    PENEMPATAN_ORDER,
    build_nama_block,
    sanitize_sheet_name,
    penempatan_sort_key,
)


def is_group_header(cell_value):
    return isinstance(cell_value, str) and "Balai Pelaksanaan Prasarana Strategis" in cell_value


def get_cell_value_with_merge(ws, row, col):
    """Ambil nilai sel, termasuk dari merge cell (top-left)."""
    cell = ws.cell(row, col)
    if cell.value is not None:
        return cell.value
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row and
                merged_range.min_col <= col <= merged_range.max_col):
            return ws.cell(merged_range.min_row, merged_range.min_col).value
    return None


def validate():
    wb_data = load_workbook("Data.xlsx")
    ws_data = wb_data.active
    data_rows = []
    for r in range(2, ws_data.max_row + 1):
        nama = ws_data.cell(r, 1).value or ""
        nip = ws_data.cell(r, 2).value or ""
        pangkat = ws_data.cell(r, 3).value or ""
        golongan = ws_data.cell(r, 4).value or ""
        tmt = ws_data.cell(r, 5).value or ""
        jabatan = ws_data.cell(r, 6).value or ""
        unit_org = ws_data.cell(r, 7).value or ""
        unit_kerja = ws_data.cell(r, 8).value or ""
        balai_penugasan = ws_data.cell(r, 10).value or ""
        if any(str(v).strip() for v in [nama, nip, jabatan, unit_org, unit_kerja, balai_penugasan]):
            data_rows.append({
                "nama": nama, "nip": nip, "pangkat": pangkat, "golongan": golongan,
                "tmt": tmt, "jabatan": jabatan, "unit_org": unit_org,
                "unit_kerja": unit_kerja, "balai_penugasan": balai_penugasan,
            })
    wb_data.close()

    print(f"Data rows: {len(data_rows)}")
    unit_orgs = sorted(set(r["unit_org"] for r in data_rows))
    used_names = {"all"}
    expected_sheets = ["all"] + [sanitize_sheet_name(u, used_names) for u in unit_orgs]

    # -----------------------------------------------------------------------
    # Template 1
    # -----------------------------------------------------------------------
    wb1 = load_workbook("Output Template 1.xlsx")
    assert wb1.sheetnames == expected_sheets, f"Template 1 sheets mismatch: {wb1.sheetnames} vs {expected_sheets}"
    errors = 0
    for sheet in wb1.sheetnames:
        ws = wb1[sheet]
        if sheet == "all":
            rows = data_rows
        else:
            unit = [u for u in unit_orgs if sanitize_sheet_name(u, {"all"}) == sheet][0]
            rows = [r for r in data_rows if r["unit_org"] == unit]

        # Build expected sequence
        groups = OrderedDict()
        for row in rows:
            key = row["balai_penugasan"]
            if key not in groups:
                groups[key] = []
            groups[key].append(row)

        expected = []
        for penempatan, group_rows in sorted(groups.items(), key=penempatan_sort_key):
            expected.append(("header", penempatan))
            for row in group_rows:
                expected.append(("data", row))

        actual = []
        for r in range(7, ws.max_row + 1):
            val = ws.cell(r, 1).value
            if is_group_header(val):
                actual.append(("header", val))
            else:
                actual.append(("data", {
                    "nama_block": ws.cell(r, 2).value or "",
                    "jabatan": ws.cell(r, 3).value or "",
                    "unit_kerja": ws.cell(r, 4).value or "",
                    "unit_org": ws.cell(r, 5).value or "",
                }))

        if len(actual) != len(expected):
            errors += 1
            print(f"Template 1 sheet {sheet}: length mismatch {len(actual)} vs {len(expected)}")
            continue

        for i, (a, e) in enumerate(zip(actual, expected)):
            if a[0] != e[0]:
                errors += 1
                continue
            if a[0] == "header":
                if a[1] != e[1]:
                    errors += 1
            else:
                row = e[1]
                expected_block = build_nama_block(
                    row["nama"], row["nip"], row["pangkat"], row["golongan"], include_tmt=True, tmt=row["tmt"]
                )
                if a[1]["nama_block"] != expected_block: errors += 1
                if a[1]["jabatan"] != row["jabatan"]: errors += 1
                if a[1]["unit_kerja"] != row["unit_kerja"]: errors += 1
                if a[1]["unit_org"] != row["unit_org"]: errors += 1
    wb1.close()
    print(f"Template 1 errors: {errors}")

    # -----------------------------------------------------------------------
    # Template 2
    # -----------------------------------------------------------------------
    wb2 = load_workbook("Output Template 2.xlsx")
    assert wb2.sheetnames == expected_sheets, f"Template 2 sheets mismatch: {wb2.sheetnames} vs {expected_sheets}"
    errors = 0
    for sheet in wb2.sheetnames:
        ws = wb2[sheet]
        if sheet == "all":
            rows = data_rows
        else:
            unit = [u for u in unit_orgs if sanitize_sheet_name(u, {"all"}) == sheet][0]
            rows = [r for r in data_rows if r["unit_org"] == unit]

        # Build expected groups: (penempatan, unit_org, unit_kerja) -> {jabatan: count}
        groups = OrderedDict()
        for row in rows:
            group_key = (row["balai_penugasan"], row["unit_org"], row["unit_kerja"])
            if group_key not in groups:
                groups[group_key] = OrderedDict()
            jabatan = row["jabatan"]
            groups[group_key][jabatan] = groups[group_key].get(jabatan, 0) + 1

        sorted_groups = sorted(groups.items(), key=lambda x: (
            penempatan_sort_key(x[0][0]), x[0][1], x[0][2]
        ))

        expected = []
        exp_no = 0
        exp_penempatan = None
        exp_grand_total = 0
        for (penempatan, unit_org, unit_kerja), jabatan_counts in sorted_groups:
            if penempatan != exp_penempatan:
                exp_no += 1
                exp_penempatan = penempatan
            total = sum(jabatan_counts.values())
            exp_grand_total += total
            for jabatan, count in jabatan_counts.items():
                expected.append((exp_no, penempatan, unit_org, unit_kerja, jabatan, count, total))
        # Baris TOTAL di bawah data: merge A:E "TOTAL", F & G = total keseluruhan
        expected.append(("TOTAL", "TOTAL", "TOTAL", "TOTAL", "", exp_grand_total, exp_grand_total))

        actual = []
        for r in range(7, ws.max_row + 1):
            actual.append((
                get_cell_value_with_merge(ws, r, 1),
                get_cell_value_with_merge(ws, r, 2) or "",
                get_cell_value_with_merge(ws, r, 3) or "",
                get_cell_value_with_merge(ws, r, 4) or "",
                ws.cell(r, 5).value or "",
                ws.cell(r, 6).value,
                get_cell_value_with_merge(ws, r, 7),
            ))

        if len(actual) != len(expected):
            errors += 1
            print(f"Template 2 sheet {sheet}: length mismatch {len(actual)} vs {len(expected)}")
            continue

        for a, e in zip(actual, expected):
            if a != e:
                errors += 1
    wb2.close()
    print(f"Template 2 errors: {errors}")

    # -----------------------------------------------------------------------
    # Template 3
    # -----------------------------------------------------------------------
    wb3 = load_workbook("Output Template 3.xlsx")
    ws3 = wb3.active
    errors = 0
    sorted_rows = sorted(data_rows, key=lambda r: penempatan_sort_key(r["balai_penugasan"]))
    for i, row in enumerate(sorted_rows):
        r = 5 + i
        expected_b = f"{row['nama']}\nNIP. {row['nip']}\n{row['jabatan']}"
        expected_c = f"{row['unit_kerja']}, {row['unit_org']}"
        if ws3.cell(r, 1).value != i + 1: errors += 1
        if ws3.cell(r, 2).value != expected_b: errors += 1
        if ws3.cell(r, 3).value != expected_c: errors += 1
        if ws3.cell(r, 4).value != row["balai_penugasan"]: errors += 1
    wb3.close()
    print(f"Template 3 errors: {errors}")

    # -----------------------------------------------------------------------
    # Template 4
    # -----------------------------------------------------------------------
    wb4 = load_workbook("Output Template 4.xlsx")
    assert wb4.sheetnames == expected_sheets, f"Template 4 sheets mismatch: {wb4.sheetnames} vs {expected_sheets}"
    errors = 0
    for sheet in wb4.sheetnames:
        ws = wb4[sheet]
        if sheet == "all":
            rows = data_rows
        else:
            unit = [u for u in unit_orgs if sanitize_sheet_name(u, {"all"}) == sheet][0]
            rows = [r for r in data_rows if r["unit_org"] == unit]

        sorted_rows = sorted(rows, key=lambda r: penempatan_sort_key(r["balai_penugasan"]))
        for i, row in enumerate(sorted_rows):
            r = 6 + i
            expected_b = build_nama_block(row["nama"], row["nip"], row["pangkat"], row["golongan"])
            if ws.cell(r, 1).value != i + 1: errors += 1
            if ws.cell(r, 2).value != expected_b: errors += 1
            if ws.cell(r, 3).value != row["jabatan"]: errors += 1
            if ws.cell(r, 4).value != row["unit_kerja"]: errors += 1
            if ws.cell(r, 5).value != row["balai_penugasan"]: errors += 1
    wb4.close()
    print(f"Template 4 errors: {errors}")

    if errors == 0:
        print("\nSemua output valid!")


if __name__ == "__main__":
    validate()

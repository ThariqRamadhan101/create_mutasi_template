import os
import re
import shutil
from collections import OrderedDict
from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Konfigurasi
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Urutan Unit Kerja Penempatan sesuai permintaan user
PENEMPATAN_ORDER = [
    "Balai Pelaksanaan Prasarana Strategis Aceh",
    "Balai Pelaksanaan Prasarana Strategis Sumatera Utara",
    "Balai Pelaksanaan Prasarana Strategis Sumatera Barat",
    "Balai Pelaksanaan Prasarana Strategis Riau",
    "Balai Pelaksanaan Prasarana Strategis Kepulauan Riau",
    "Balai Pelaksanaan Prasarana Strategis Jambi",
    "Balai Pelaksanaan Prasarana Strategis Bengkulu",
    "Balai Pelaksanaan Prasarana Strategis Sumatera Selatan",
    "Balai Pelaksanaan Prasarana Strategis Kepulauan Bangka Belitung",
    "Balai Pelaksanaan Prasarana Strategis Lampung",
    "Balai Pelaksanaan Prasarana Strategis DKI Jakarta",
    "Balai Pelaksanaan Prasarana Strategis Jawa Barat",
    "Balai Pelaksanaan Prasarana Strategis Banten",
    "Balai Pelaksanaan Prasarana Strategis Jawa Tengah",
    "Balai Pelaksanaan Prasarana Strategis DI Yogyakarta",
    "Balai Pelaksanaan Prasarana Strategis Jawa Timur",
    "Balai Pelaksanaan Prasarana Strategis Bali",
    "Balai Pelaksanaan Prasarana Strategis Nusa Tenggara Barat",
    "Balai Pelaksanaan Prasarana Strategis Nusa Tenggara Timur",
    "Balai Pelaksanaan Prasarana Strategis Kalimantan Barat",
    "Balai Pelaksanaan Prasarana Strategis Kalimantan Tengah",
    "Balai Pelaksanaan Prasarana Strategis Kalimantan Selatan",
    "Balai Pelaksanaan Prasarana Strategis Kalimantan Timur",
    "Balai Pelaksanaan Prasarana Strategis Kalimantan Utara",
    "Balai Pelaksanaan Prasarana Strategis Sulawesi Utara & Gorontalo",
    "Balai Pelaksanaan Prasarana Strategis Sulawesi Tengah",
    "Balai Pelaksanaan Prasarana Strategis Sulawesi Selatan",
    "Balai Pelaksanaan Prasarana Strategis Sulawesi Tenggara",
    "Balai Pelaksanaan Prasarana Strategis Sulawesi Barat",
    "Balai Pelaksanaan Prasarana Strategis Maluku",
    "Balai Pelaksanaan Prasarana Strategis Maluku Utara",
    "Balai Pelaksanaan Prasarana Strategis Papua",
    "Balai Pelaksanaan Prasarana Strategis Papua Barat",
]

DATA_FILE = os.path.join(BASE_DIR, "Data.xlsx")
TEMPLATE_FILES = {
    "template1": os.path.join(BASE_DIR, "Template 1.xlsx"),
    "template2": os.path.join(BASE_DIR, "Template 2.xlsx"),
    "template3": os.path.join(BASE_DIR, "Template 3.xlsx"),
    "template4": os.path.join(BASE_DIR, "Template 4.xlsx"),
}
OUTPUT_FILES = {
    "template1": os.path.join(BASE_DIR, "Output Template 1.xlsx"),
    "template2": os.path.join(BASE_DIR, "Output Template 2.xlsx"),
    "template3": os.path.join(BASE_DIR, "Output Template 3.xlsx"),
    "template4": os.path.join(BASE_DIR, "Output Template 4.xlsx"),
}

BULAN_INDONESIA = {
    "01": "Januari", "02": "Februari", "03": "Maret", "04": "April",
    "05": "Mei", "06": "Juni", "07": "Juli", "08": "Agustus",
    "09": "September", "10": "Oktober", "11": "November", "12": "Desember",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def sanitize_sheet_name(name, used_names=None):
    r"""
    Bersihkan nama sheet agar valid di Excel.
    Excel: maksimal 31 karakter, tidak boleh mengandung \ / ? * [ ]
    """
    used_names = used_names or set()
    name = re.sub(r'[\\/*?:\[\]]', '_', str(name))
    name = name.strip()
    base = name[:31]
    candidate = base
    counter = 1
    while candidate in used_names or candidate.lower() == "all":
        suffix = f"_{counter}"
        candidate = base[:31 - len(suffix)] + suffix
        counter += 1
    return candidate


def penempatan_sort_key(item):
    """Key untuk mengurutkan grup berdasarkan urutan yang ditentukan."""
    name = item[0] if isinstance(item, tuple) else item
    if name in PENEMPATAN_ORDER:
        return (PENEMPATAN_ORDER.index(name), name)
    return (len(PENEMPATAN_ORDER), name)


def format_tmt(tmt_value):
    """Format TMT dari DD-MM-YYYY menjadi 'TMT DD Bulan YYYY'."""
    if not tmt_value:
        return "TMT -"
    tmt_str = str(tmt_value).strip()
    # Coba parsing DD-MM-YYYY atau DD/MM/YYYY
    for sep in ["-", "/"]:
        if sep in tmt_str:
            parts = tmt_str.split(sep)
            if len(parts) == 3:
                dd, mm, yyyy = parts
                if len(yyyy) == 2:
                    yyyy = "20" + yyyy
                bulan = BULAN_INDONESIA.get(mm, mm)
                return f"TMT {dd} {bulan} {yyyy}"
    return f"TMT {tmt_str}"


def build_nama_block(nama, nip, pangkat, golongan, include_tmt=False, tmt=None):
    """Susun blok Nama - NIP - Pang/Gol (opsional TMT)."""
    lines = [nama]
    if nip:
        lines.append(f"NIP. {nip}")
    pangkat_gol = ""
    if pangkat and golongan:
        pangkat_gol = f"{pangkat} ({golongan})"
    elif pangkat:
        pangkat_gol = pangkat
    elif golongan:
        pangkat_gol = f"({golongan})"
    if pangkat_gol:
        lines.append(pangkat_gol)
    if include_tmt and tmt is not None:
        lines.append(format_tmt(tmt))
    return "\n".join(lines)


def read_data_rows(ws):
    """Baca baris data dari Data.xlsx (struktur baru) mulai baris ke-2."""
    rows = []
    for r in range(2, ws.max_row + 1):
        nama = ws.cell(r, 1).value or ""
        nip = ws.cell(r, 2).value or ""
        pangkat = ws.cell(r, 3).value or ""
        golongan = ws.cell(r, 4).value or ""
        tmt = ws.cell(r, 5).value or ""
        jabatan = ws.cell(r, 6).value or ""
        unit_org = ws.cell(r, 7).value or ""
        unit_kerja = ws.cell(r, 8).value or ""
        propinsi = ws.cell(r, 9).value or ""
        balai_penugasan = ws.cell(r, 10).value or ""

        if any(str(v).strip() for v in [nama, nip, jabatan, unit_org, unit_kerja, balai_penugasan]):
            rows.append({
                "nama": nama,
                "nip": nip,
                "pangkat": pangkat,
                "golongan": golongan,
                "tmt": tmt,
                "jabatan": jabatan,
                "unit_org": unit_org,
                "unit_kerja": unit_kerja,
                "propinsi": propinsi,
                "balai_penugasan": balai_penugasan,
            })
    return rows


def capture_row_style(ws, row, max_col):
    """Ambil style dari setiap sel pada baris tertentu."""
    styles = []
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        styles.append({
            "font": copy(cell.font),
            "border": copy(cell.border),
            "fill": copy(cell.fill),
            "number_format": copy(cell.number_format),
            "protection": copy(cell.protection),
            "alignment": copy(cell.alignment),
        })
    return styles


def apply_row_style(ws, row, styles, height):
    """Terapkan style yang sudah di-capture ke baris tertentu."""
    if height is not None:
        ws.row_dimensions[row].height = height
    for c, style in enumerate(styles, start=1):
        cell = ws.cell(row, c)
        cell.font = style["font"]
        cell.border = style["border"]
        cell.fill = style["fill"]
        cell.number_format = style["number_format"]
        cell.protection = style["protection"]
        cell.alignment = style["alignment"]


def group_by_penempatan(rows):
    """Kelompokkan baris berdasarkan Unit Kerja Penempatan sesuai urutan tetap."""
    groups = OrderedDict()
    for row in rows:
        key = row["balai_penugasan"]
        if key not in groups:
            groups[key] = []
        groups[key].append(row)
    return sorted(groups.items(), key=penempatan_sort_key)


# ---------------------------------------------------------------------------
# Template 1
# ---------------------------------------------------------------------------
def build_template1_sheet(ws, rows):
    """Susun worksheet Template 1 dengan grup per Unit Kerja Penempatan."""
    max_col = ws.max_column

    # Capture style baris grup (baris 7) dan baris data (baris 8)
    group_header_styles = capture_row_style(ws, 7, max_col)
    data_row_styles = capture_row_style(ws, 8, max_col)
    group_header_height = ws.row_dimensions[7].height
    data_row_height = ws.row_dimensions[8].height

    # Hapus semua baris dari baris 7 ke bawah, lalu susun ulang
    ws.delete_rows(7, ws.max_row - 6)

    current_row = 7
    no = 1
    for penempatan, group_rows in group_by_penempatan(rows):
        # Baris header grup (merge A:E)
        ws.insert_rows(current_row)
        apply_row_style(ws, current_row, group_header_styles, group_header_height)
        ws.cell(current_row, 1).value = penempatan
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        current_row += 1

        # Baris data
        for row in group_rows:
            ws.insert_rows(current_row)
            apply_row_style(ws, current_row, data_row_styles, data_row_height)
            ws.cell(current_row, 1).value = no
            ws.cell(current_row, 2).value = build_nama_block(
                row["nama"], row["nip"], row["pangkat"], row["golongan"], include_tmt=True, tmt=row["tmt"]
            )
            ws.cell(current_row, 3).value = row["jabatan"]
            ws.cell(current_row, 4).value = row["unit_kerja"]
            ws.cell(current_row, 5).value = row["unit_org"]
            current_row += 1
            no += 1


def generate_template1(data_rows):
    src = TEMPLATE_FILES["template1"]
    dst = OUTPUT_FILES["template1"]
    shutil.copy2(src, dst)

    wb = load_workbook(dst)
    ws_template = wb.active

    # Sheet "all"
    ws_all = wb.copy_worksheet(ws_template)
    ws_all.title = "all"
    build_template1_sheet(ws_all, data_rows)

    # Sheet per Unit Organisasi Asal
    unit_orgs = sorted(set(row["unit_org"] for row in data_rows))
    used_names = {"all"}
    for unit_org in unit_orgs:
        sheet_name = sanitize_sheet_name(unit_org, used_names)
        used_names.add(sheet_name)
        ws_unit = wb.copy_worksheet(ws_template)
        ws_unit.title = sheet_name
        filtered_rows = [row for row in data_rows if row["unit_org"] == unit_org]
        build_template1_sheet(ws_unit, filtered_rows)

    wb.remove(ws_template)
    wb.save(dst)
    wb.close()
    print(f"[OK] {dst} - {len(data_rows)} baris data, {len(unit_orgs)} unit organisasi")


# ---------------------------------------------------------------------------
# Template 2
# ---------------------------------------------------------------------------
def build_template2_sheet(ws, rows):
    """Susun worksheet Template 2 dengan merge hierarkis per kolom."""
    max_col = ws.max_column

    # Capture style baris data (baris 7) dan tinggi baris
    data_row_styles = capture_row_style(ws, 7, max_col)
    data_row_height = ws.row_dimensions[7].height

    # Hapus merge cells di area data agar tidak mengganggu baris baru
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= 7:
            ws.unmerge_cells(str(merged_range))

    # Hapus semua baris dari baris 7 ke bawah, lalu susun ulang
    ws.delete_rows(7, ws.max_row - 6)

    # Agregasi: grup -> {jabatan: count}
    groups = OrderedDict()
    for row in rows:
        group_key = (row["balai_penugasan"], row["unit_org"], row["unit_kerja"])
        if group_key not in groups:
            groups[group_key] = OrderedDict()
        jabatan = row["jabatan"]
        groups[group_key][jabatan] = groups[group_key].get(jabatan, 0) + 1

    # Urutkan grup berdasarkan urutan penempatan, lalu unit_org, unit_kerja
    sorted_groups = sorted(groups.items(), key=lambda x: (
        penempatan_sort_key(x[0][0]), x[0][1], x[0][2]
    ))

    # Bangun semua baris dan catat batas setiap grup
    current_row = 7
    current_no = 0
    current_penempatan = None
    group_boundaries = []  # (start_row, end_row, penempatan, unit_org, unit_kerja)
    for (penempatan, unit_org, unit_kerja), jabatan_counts in sorted_groups:
        if penempatan != current_penempatan:
            current_no += 1
            current_penempatan = penempatan

        n_rows = len(jabatan_counts)
        total = sum(jabatan_counts.values())

        ws.insert_rows(current_row, n_rows)
        for offset in range(n_rows):
            apply_row_style(ws, current_row + offset, data_row_styles, data_row_height)

        # Isi nilai
        ws.cell(current_row, 1).value = current_no
        ws.cell(current_row, 2).value = penempatan
        ws.cell(current_row, 3).value = unit_org
        ws.cell(current_row, 4).value = unit_kerja
        ws.cell(current_row, 7).value = total

        # Isi Jabatan dan Jumlah per baris
        for offset, (jabatan, count) in enumerate(jabatan_counts.items()):
            ws.cell(current_row + offset, 5).value = jabatan
            ws.cell(current_row + offset, 6).value = count

        group_boundaries.append((current_row, current_row + n_rows - 1, penempatan, unit_org, unit_kerja))
        current_row += n_rows

    # Merge kolom D (UPT Asal) dan G (Total) dalam setiap grup
    for start_row, end_row, _, _, _ in group_boundaries:
        if start_row < end_row:
            ws.merge_cells(start_row=start_row, start_column=4, end_row=end_row, end_column=4)
            ws.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)

    # Merge kolom C (Unit Organisasi Asal) selama (penempatan, unit_org) sama
    i = 0
    while i < len(group_boundaries):
        start_row = group_boundaries[i][0]
        penempatan = group_boundaries[i][2]
        unit_org = group_boundaries[i][3]
        j = i + 1
        while j < len(group_boundaries) and group_boundaries[j][2] == penempatan and group_boundaries[j][3] == unit_org:
            j += 1
        end_row = group_boundaries[j - 1][1]
        if start_row < end_row:
            ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
        i = j

    # Merge kolom A (No) dan B (Unit Kerja Penempatan) selama penempatan sama
    i = 0
    while i < len(group_boundaries):
        start_row = group_boundaries[i][0]
        penempatan = group_boundaries[i][2]
        j = i + 1
        while j < len(group_boundaries) and group_boundaries[j][2] == penempatan:
            j += 1
        end_row = group_boundaries[j - 1][1]
        if start_row < end_row:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
        i = j


def generate_template2(data_rows):
    src = TEMPLATE_FILES["template2"]
    dst = OUTPUT_FILES["template2"]
    shutil.copy2(src, dst)

    wb = load_workbook(dst)
    ws_template = wb.active

    # Sheet "all"
    ws_all = wb.copy_worksheet(ws_template)
    ws_all.title = "all"
    build_template2_sheet(ws_all, data_rows)

    # Sheet per Unit Organisasi Asal
    unit_orgs = sorted(set(row["unit_org"] for row in data_rows))
    used_names = {"all"}
    for unit_org in unit_orgs:
        sheet_name = sanitize_sheet_name(unit_org, used_names)
        used_names.add(sheet_name)
        ws_unit = wb.copy_worksheet(ws_template)
        ws_unit.title = sheet_name
        filtered_rows = [row for row in data_rows if row["unit_org"] == unit_org]
        build_template2_sheet(ws_unit, filtered_rows)

    wb.remove(ws_template)
    wb.save(dst)
    wb.close()
    print(f"[OK] {dst} - rekap total, {len(unit_orgs)} unit organisasi")


# ---------------------------------------------------------------------------
# Template 3
# ---------------------------------------------------------------------------
def generate_template3(data_rows):
    src = TEMPLATE_FILES["template3"]
    dst = OUTPUT_FILES["template3"]
    shutil.copy2(src, dst)

    wb = load_workbook(dst)
    ws = wb.active

    # Data dimulai baris 5
    start_row = 5

    # Urutkan data berdasarkan urutan penempatan
    sorted_rows = sorted(data_rows, key=lambda r: penempatan_sort_key(r["balai_penugasan"]))

    for idx, row in enumerate(sorted_rows, start=1):
        r = start_row + idx - 1

        # B: Nama - NIP - Jabatan
        nama_nip = f"{row['nama']}\nNIP. {row['nip']}"
        nama_jabatan = f"{nama_nip}\n{row['jabatan']}"
        # C: Unit Kerja, Unit Organisasi Asal
        unit_asal = f"{row['unit_kerja']}, {row['unit_org']}"

        ws.cell(r, 1).value = idx
        ws.cell(r, 2).value = nama_jabatan
        ws.cell(r, 3).value = unit_asal
        ws.cell(r, 4).value = row["balai_penugasan"]

    wb.save(dst)
    wb.close()
    print(f"[OK] {dst} - {len(data_rows)} baris data")


# ---------------------------------------------------------------------------
# Template 4
# ---------------------------------------------------------------------------
def build_template4_sheet(ws, rows, unit_org_title=None):
    """Susun worksheet Template 4."""
    max_col = ws.max_column

    # Update judul sesuai unit organisasi asal
    title_cell = ws.cell(2, 1)
    if unit_org_title:
        title_cell.value = f"Daftar Usulan Mutasi PNS\n{unit_org_title}"
    else:
        title_cell.value = "Daftar Usulan Mutasi PNS\nSEMUA UNIT ORGANISASI ASAL"

    # Capture style baris data (baris 6)
    data_row_styles = capture_row_style(ws, 6, max_col)
    data_row_height = ws.row_dimensions[6].height

    # Hapus semua baris data dari baris 6 ke bawah, lalu susun ulang
    ws.delete_rows(6, ws.max_row - 5)

    # Urutkan data berdasarkan urutan penempatan
    sorted_rows = sorted(rows, key=lambda r: penempatan_sort_key(r["balai_penugasan"]))

    current_row = 6
    for idx, row in enumerate(sorted_rows, start=1):
        ws.insert_rows(current_row)
        apply_row_style(ws, current_row, data_row_styles, data_row_height)
        ws.cell(current_row, 1).value = idx
        ws.cell(current_row, 2).value = build_nama_block(
            row["nama"], row["nip"], row["pangkat"], row["golongan"]
        )
        ws.cell(current_row, 3).value = row["jabatan"]
        ws.cell(current_row, 4).value = row["unit_kerja"]
        ws.cell(current_row, 5).value = row["balai_penugasan"]
        current_row += 1


def generate_template4(data_rows):
    src = TEMPLATE_FILES["template4"]
    dst = OUTPUT_FILES["template4"]
    shutil.copy2(src, dst)

    wb = load_workbook(dst)
    ws_template = wb.active

    # Sheet "all"
    ws_all = wb.copy_worksheet(ws_template)
    ws_all.title = "all"
    build_template4_sheet(ws_all, data_rows)

    # Sheet per Unit Organisasi Asal
    unit_orgs = sorted(set(row["unit_org"] for row in data_rows))
    used_names = {"all"}
    for unit_org in unit_orgs:
        sheet_name = sanitize_sheet_name(unit_org, used_names)
        used_names.add(sheet_name)
        ws_unit = wb.copy_worksheet(ws_template)
        ws_unit.title = sheet_name
        filtered_rows = [row for row in data_rows if row["unit_org"] == unit_org]
        build_template4_sheet(ws_unit, filtered_rows, unit_org_title=unit_org)

    wb.remove(ws_template)
    wb.save(dst)
    wb.close()
    print(f"[OK] {dst} - {len(data_rows)} baris data, {len(unit_orgs)} unit organisasi")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    if not os.path.exists(DATA_FILE):
        raise FileNotFoundError(f"File data tidak ditemukan: {DATA_FILE}")
    for key, path in TEMPLATE_FILES.items():
        if not os.path.exists(path):
            raise FileNotFoundError(f"File template tidak ditemukan: {path}")

    wb_data = load_workbook(DATA_FILE, data_only=False)
    ws_data = wb_data.active
    data_rows = read_data_rows(ws_data)
    wb_data.close()

    print(f"Membaca {len(data_rows)} baris data dari Data.xlsx")

    generate_template1(data_rows)
    generate_template2(data_rows)
    generate_template3(data_rows)
    generate_template4(data_rows)

    print("\nSelesai. File output telah dibuat di folder yang sama.")


if __name__ == "__main__":
    main()
